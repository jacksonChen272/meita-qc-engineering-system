from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
QC_FIXTURE = ROOT / "input" / "qc-template.docx"
RND_FIXTURE = next(iter(sorted((ROOT / "input").glob("*-AF_*.docx"))), None)

from diff import align_display_value, build_diff, values_equal
from mapping import map_parameters
from parsers import parse_legacy_metadata, parse_qc_template, parse_rnd_document
from render import build_preview_model, validate_pagination


class NormalizeTests(unittest.TestCase):
    def test_format_variants_are_equal(self) -> None:
        self.assertTrue(values_equal("81 ± 3 °C", "81±3℃"))
        self.assertTrue(values_equal("81 ±3 C", "81±3°C"))
        self.assertTrue(values_equal("35~60", "35–60 mg/100 g"))
        self.assertTrue(values_equal("13~20", "13–20 °C"))

    def test_different_values_stay_different(self) -> None:
        self.assertFalse(values_equal("121°C", "121.5±0.5°C"))
        self.assertFalse(values_equal("10 min", "15 min"))
        self.assertFalse(values_equal("254±4 g", "252–263 g"))

    def test_display_follows_unitless_master_and_keeps_protection_marker(self) -> None:
        self.assertEqual(align_display_value("200±25", "200 ± 25 bar"), "200 ± 25")
        self.assertEqual(align_display_value("23(CCP)", "17 分鐘"), "17(CCP)")
        self.assertEqual(align_display_value("4(OPRP)", "4 rpm"), "4(OPRP)")


class LegacyMetadataTests(unittest.TestCase):
    def test_old_qc_header_without_cover_page_is_supported(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "legacy.docx"
            document = Document()
            table = document.add_table(rows=3, cols=7)
            values = [
                ["力增 洗腎配方(杏仁)\nQC工程圖", "適 用 範 圍", "生 產 二 課", "類 別", "品質管制標準書", "編 號", "02-0200-051"],
                ["", "工程圖記號定義", "", "制 訂", "108年05月29日", "發佈", "108年05月31日"],
                ["", "", "", "修 訂", "第6次修訂：113年06月06日", "", ""],
            ]
            for row_index, row in enumerate(values):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
            document.save(path)
            metadata = parse_legacy_metadata(path)

        self.assertEqual(metadata["documentNumber"], "02-0200-051")
        self.assertEqual(metadata["establishedDate"], "108.05.29")
        self.assertEqual(metadata["unit"], "生產二課")
        self.assertEqual(metadata["latestRevision"], "6.0")
        self.assertEqual(metadata["revisionHistory"][0]["revisionDate"], "113.06.06")


@unittest.skipUnless(
    QC_FIXTURE.exists() and RND_FIXTURE is not None,
    "本機整合測試需要 QC 母版與保留受控檔名的外來標準",
)
class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.qc = parse_qc_template(QC_FIXTURE)
        cls.rnd = parse_rnd_document(RND_FIXTURE)
        cls.mapping = map_parameters(cls.rnd, cls.qc)
        cls.diff = build_diff(cls.qc, cls.rnd, cls.mapping)

    def test_actual_document_counts(self) -> None:
        self.assertEqual(self.qc["processStepCount"], 55)
        self.assertEqual(self.qc["controlItemCount"], 116)
        self.assertEqual(self.rnd["parameterCount"], 82)

    def test_product_metadata_is_parsed(self) -> None:
        self.assertEqual(self.rnd["product"]["version"], "1.10")
        self.assertEqual(self.rnd["product"]["date"], "2025.11.03")
        self.assertIn("洗腎配方(杏仁)", self.rnd["product"]["name"])

    def test_locked_fields_are_exactly_explicit_special_food_fields(self) -> None:
        locked = [record for record in self.diff["records"] if record["status"] == "LOCKED"]
        self.assertEqual(len(locked), 4)
        self.assertEqual({record["controlItem"].split("(")[0] for record in locked}, {"內容量", "脂肪", "容積", "水份"})

    def test_filling_weight_never_overwrites_content_weight(self) -> None:
        filling = next(parameter for parameter in self.rnd["parameters"] if parameter["category"] == "filling" and parameter["name"] == "填充量")
        mapped = [item for item in self.mapping["mappings"] if item["rndParameterId"] == filling["id"]]
        self.assertTrue(any(item["targetControlItemName"].startswith("全重") and not item["forceReview"] for item in mapped))
        self.assertFalse(any(item["targetControlItemName"].startswith("內容量") and not item["forceReview"] for item in mapped))

    def test_invalid_upper_retort_temperature_requires_review(self) -> None:
        record = next(record for record in self.diff["records"] if record["process"] == "殺菌" and "上釜" in record["controlItem"])
        self.assertEqual(record["rndValue"], "X")
        self.assertEqual(record["status"], "REVIEW_REQUIRED")

    def test_superscript_and_filler_lines_are_preserved_correctly(self) -> None:
        raw_water = next(step for step in self.qc["processSteps"] if step["name"] == "原水存放")
        total_count = next(control for control in raw_water["control_items"] if control["name"].startswith("總生菌數"))
        self.assertEqual(total_count["standard"], "≦10²")
        for page in self.qc["layout"]["pages"]:
            for row in page["rows"]:
                lines = row["fields"]["machine"]["lines"]
                if lines and all(set(line) <= {"-"} for line in lines):
                    self.assertEqual(len(lines), 1)

    def test_confirmed_text_and_sampling_location_corrections(self) -> None:
        all_control_text = "\n".join(row["fields"]["controlItem"]["text"] for page in self.qc["layout"]["pages"] for row in page["rows"])
        self.assertIn("T.W.C、罐徑", all_control_text)
        self.assertNotIn("T..W..C", all_control_text)
        expected = {
            "空罐儲放": "原物料儲放區",
            "罐蓋檢驗": "原物料儲放區",
            "罐蓋儲放": "原物料儲放區",
            "真空檢測打檢": "包裝區",
            "成品開罐": "包裝區",
        }
        by_name = {step["name"]: step for step in self.qc["processSteps"]}
        for name, location in expected.items():
            self.assertEqual(by_name[name]["sampling_locations"], [location])
        water_rinse = next(item for item in by_name["空罐噴洗"]["control_items"] if "水溫" in item["name"])
        self.assertEqual(water_rinse["standard"], "≧70")
        self.assertEqual(by_name["包裝"]["operation_standards"], ["02-0203-042"])
        self.assertEqual(by_name["出貨"]["sampling_locations"], ["出貨區"])

    def test_seam_retort_packaging_and_end_flow_match_confirmed_reference(self) -> None:
        by_name = {step["name"]: step for step in self.qc["processSteps"]}
        sealing = by_name["封蓋"]["control_items"][0]
        seam = by_name["捲封檢查"]
        self.assertEqual(sealing["item_details"], "切罐\n凹罐\n層狀\n漏罐\n舌狀")
        self.assertEqual([item["item_details"] for item in seam["control_items"]], ["切罐\n凹罐\n層狀\n漏罐\n舌狀", "OL%\nWR\nT.W.C.BH\nCH.OL"])
        seam_pages = {
            page["id"]
            for page in self.qc["layout"]["pages"]
            for row in page["rows"]
            if row["processStepId"] == seam["id"]
        }
        self.assertEqual(len(seam_pages), 1)
        retort = by_name["殺菌"]
        self.assertEqual(retort["machines"], ["旋轉式熱水淋灑臥式高溫高壓殺菌釜"])
        self.assertEqual([item["name"] for item in retort["control_items"]], ["昇溫時間(min)", "溫度(℃)", "時間(min)", "轉速(rpm)"])
        self.assertEqual([item["standard"] for item in retort["control_items"]], ["17(CCP)", "121.5±0.5(CCP)", "15(CCP)", "4(CCP)"])
        self.assertEqual(by_name["包裝檢查"]["operation_standards"], ["10-0203-004"])
        self.assertEqual([item["name"].replace(" ", "").replace("/", "") for item in by_name["包裝檢查"]["control_items"]], ["紙箱缺點", "外觀不良率"])
        self.assertEqual(by_name["裝箱檢查"]["operation_standards"], ["10-0203-004"])
        self.assertEqual([item["name"] for item in by_name["裝箱檢查"]["control_items"]], ["捲封不良率", "嚴重缺點"])
        self.assertTrue(all(item["chart"] == "成品包裝品質檢查表" for name in ("包裝檢查", "裝箱檢查") for item in by_name[name]["control_items"]))
        self.assertTrue(all("包裝場" not in location for step in self.qc["processSteps"] for location in step["sampling_locations"]))
        case_check = by_name["整箱打檢"]
        self.assertEqual(case_check["default_sampling_quantities"], ["5批/次"])
        self.assertEqual(len(case_check["control_items"]), 2)
        case_check_pages = {
            page["id"]
            for page in self.qc["layout"]["pages"]
            for row in page["rows"]
            if row["processStepId"] == case_check["id"]
        }
        self.assertEqual(len(case_check_pages), 1)
        self.assertEqual(case_check["flow_symbol"], "inspection")
        self.assertEqual(by_name["出貨"]["flow_symbol"], "end")

    def test_protected_master_markers_are_never_lost(self) -> None:
        records = [record for record in self.diff["records"] if "(CCP)" in record["templateValue"] or "(OPRP)" in record["templateValue"]]
        self.assertTrue(records)
        self.assertTrue(all(not record["rndValue"] or "(CCP)" in record["rndValue"] or "(OPRP)" in record["rndValue"] for record in records))

    def test_incubation_uses_master_description_and_finished_product_values(self) -> None:
        equipment = next(record for record in self.diff["records"] if record["process"] == "保溫試驗" and record["controlItem"].startswith("保溫試驗-37"))
        self.assertEqual(equipment["proposedValue"], equipment["templateValue"])
        self.assertEqual(equipment["decision"], "KEEP_TEMPLATE")
        finished = {record["controlItem"]: record["proposedValue"] for record in self.diff["records"] if record["process"] == "成品開罐"}
        incubation = {record["controlItem"]: record["proposedValue"] for record in self.diff["records"] if record["process"] == "保溫試驗"}
        self.assertEqual(incubation["真空度(cmHg)"], finished["真空度(cmHg)"])
        self.assertEqual(incubation["pH值"], finished["pH值"])
        self.assertEqual(incubation["雜質"], finished["雜質"])

    def test_unit_only_context_candidates_are_hidden(self) -> None:
        unit_only = [
            record
            for record in self.diff["records"]
            if record["classification"] == "CONDITIONAL"
            and values_equal(record["templateValue"], record["rndValue"])
        ]
        self.assertEqual(unit_only, [])

    def test_process_steps_do_not_cross_preview_pages(self) -> None:
        self.assertTrue(validate_pagination(self.qc)["valid"])

    def test_preview_repeats_document_header_on_every_page(self) -> None:
        preview = build_preview_model(self.qc, self.diff, self.rnd)
        self.assertEqual(len(preview["pages"]), 7)
        self.assertIn("洗腎配方(杏仁)", preview["documentHeader"]["title"])
        self.assertEqual(preview["documentHeader"]["category"], "品質管制標準書")
        self.assertEqual(preview["documentHeader"]["scope"], "生產二課")


if __name__ == "__main__":
    unittest.main()
